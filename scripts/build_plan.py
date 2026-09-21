#!/usr/bin/env python3
import argparse, json, os
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

FONT_CN = "Noto Sans CJK SC"
FONT_EN = "Arial"
BORDER = "D9DDE3"
HEADER_FILL = "F1F3F5"
HIGHLIGHT = "FFF59D"
LINK_BLUE = "0563C1"
TEXT = "20252B"
MUTED = "6B7280"


def set_run_font(run, size=11, bold=False, color=TEXT, font=FONT_CN):
    run.font.name = font
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=BORDER, size='4'):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None:
            node = OxmlElement('w:'+m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_keep_with_next(paragraph, keep=True):
    pPr = paragraph._p.get_or_add_pPr()
    node = pPr.find(qn('w:keepNext'))
    if keep and node is None:
        node = OxmlElement('w:keepNext')
        pPr.append(node)
    elif not keep and node is not None:
        pPr.remove(node)


def add_bottom_rule(paragraph, color="D9DDE3", size="4"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is None:
        pBdr = OxmlElement('w:pBdr')
        pPr.append(pBdr)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), '8')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)


def add_hyperlink(paragraph, text, url, font_size=11):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), LINK_BLUE); rPr.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    rFonts = OxmlElement('w:rFonts'); rFonts.set(qn('w:ascii'), FONT_CN); rFonts.set(qn('w:eastAsia'), FONT_CN); rPr.append(rFonts)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(font_size*2))); rPr.append(sz)
    new_run.append(rPr)
    t = OxmlElement('w:t'); t.text = text; new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_text(paragraph, text, size=11, bold=False, color=TEXT):
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    return run


def format_paragraph(p, before=0, after=0, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_title(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, after=18, line=1.15)
    add_text(p, text, size=21, bold=True)
    return p


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, before=2, after=8)
    add_text(p, text, size=17, bold=True)
    return p


def add_section_title(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, before=18, after=14)
    set_keep_with_next(p, True)
    add_text(p, text, size=15.5, bold=True)
    return p


def add_small_heading(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, before=14, after=10)
    set_keep_with_next(p, True)
    add_text(p, text, size=14, bold=True)
    return p


def add_rule(doc):
    p = doc.add_paragraph()
    format_paragraph(p, before=4, after=8)
    add_bottom_rule(p)
    return p


def set_table_layout(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_border(cell)
            set_cell_margins(cell)
            for p in cell.paragraphs:
                format_paragraph(p, after=0, line=1.25)


def add_background_table(doc, bg):
    table = doc.add_table(rows=7, cols=2)
    set_table_layout(table)
    widths = [Cm(8.4), Cm(8.4)]
    for row in table.rows:
        for i,w in enumerate(widths):
            row.cells[i].width = w
    shade_cell(table.cell(0,0), HEADER_FILL); shade_cell(table.cell(0,1), HEADER_FILL)
    add_text(table.cell(0,1).paragraphs[0], '个人信息', size=11, bold=True)
    fields = [
        ('院校背景', bg.get('school_background','')),
        ('GPA', bg.get('gpa','待补充')),
        ('语言成绩', bg.get('language','待补充')),
        ('申请专业', bg.get('target_major','待补充')),
        ('目标国家/地区', bg.get('regions_text','待补充')),
        ('软背景', bg.get('soft_background','待补充')),
    ]
    for r,(k,v) in enumerate(fields, start=1):
        add_text(table.cell(r,0).paragraphs[0], k, size=11, bold=True)
        add_text(table.cell(r,1).paragraphs[0], v, size=11)
    return table


def add_programs_table(doc, programs, include_academic=True):
    headers = ['学校','项目名称','QS排名'] + (['成绩要求'] if include_academic else []) + ['雅思要求','官网链接']
    cols = len(headers)
    table = doc.add_table(rows=1, cols=cols)
    set_table_layout(table)
    set_repeat_table_header(table.rows[0])
    for i,h in enumerate(headers):
        shade_cell(table.cell(0,i), HEADER_FILL)
        add_text(table.cell(0,i).paragraphs[0], h, size=10.5, bold=True)
    if include_academic:
        widths = [Cm(4.4), Cm(4.1), Cm(1.55), Cm(1.75), Cm(1.75), Cm(3.0)]
    else:
        widths = [Cm(4.5), Cm(4.4), Cm(1.7), Cm(2.0), Cm(3.6)]
    for row in table.rows:
        for i,w in enumerate(widths): row.cells[i].width = w
    for item in programs:
        row = table.add_row()
        for i,w in enumerate(widths): row.cells[i].width = w
        school = f"{item.get('school_en','')} ({item.get('school_abbr','')}) {item.get('school_zh','')}".replace(' () ',' ')
        add_text(row.cells[0].paragraphs[0], school, size=10.5)
        p = row.cells[1].paragraphs[0]
        run = add_text(p, f"{item.get('program_en','')} {item.get('program_zh','')}", size=10.5)
        if item.get('highlight'):
            shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), HIGHLIGHT)
            run._r.get_or_add_rPr().append(shd)
        add_text(row.cells[2].paragraphs[0], item.get('qs_rank',''), size=10.5)
        idx=3
        if include_academic:
            add_text(row.cells[idx].paragraphs[0], item.get('academic_requirement',''), size=10.5); idx += 1
        add_text(row.cells[idx].paragraphs[0], item.get('ielts_requirement',''), size=10.5); idx += 1
        purl = row.cells[idx].paragraphs[0]
        if item.get('official_url'):
            add_hyperlink(purl, '链接', item['official_url'], font_size=10.5)
        else:
            add_text(purl, '待补充', size=10.5, color=MUTED)
    return table


def add_timeline_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    set_table_layout(table)
    widths = [Cm(8.4), Cm(8.4)]
    for i,w in enumerate(widths): table.rows[0].cells[i].width = w
    set_repeat_table_header(table.rows[0])
    for i,h in enumerate(['时间','任务']):
        shade_cell(table.cell(0,i), HEADER_FILL)
        add_text(table.cell(0,i).paragraphs[0], h, size=10.5, bold=True)
    for item in rows:
        row = table.add_row()
        for i,w in enumerate(widths): row.cells[i].width = w
        add_text(row.cells[0].paragraphs[0], item.get('time',''), size=10.5)
        p = row.cells[1].paragraphs[0]
        task = item.get('task','')
        if 'https://www.yushiconsulting.cn' in task:
            left, _, right = task.partition('https://www.yushiconsulting.cn')
            add_text(p, left, size=10.5)
            add_hyperlink(p, 'https://www.yushiconsulting.cn', 'https://www.yushiconsulting.cn/', 10.5)
            add_text(p, right, size=10.5)
        else:
            add_text(p, task, size=10.5)
    return table


def build(data, output_path, asset_dir):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.65); sec.bottom_margin = Cm(1.65)
    sec.left_margin = Cm(1.65); sec.right_margin = Cm(1.65)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = FONT_CN; normal._element.rPr.rFonts.set(qn('w:eastAsia'), FONT_CN); normal.font.size = Pt(11)

    regions = data.get('regions') or []
    region_text = '+'.join(regions) if regions else '目标地区'
    title = f"{data.get('date','').replace('-','')}-首咨选校规划 - {data.get('student_name','XX同学')} - {data.get('intake_short','')}（{region_text}）"
    add_title(doc, title)
    add_subtitle(doc, '留学选校规划')
    p = doc.add_paragraph(); format_paragraph(p, after=8)
    add_text(p, '│ 日期： ', size=10.5, color=MUTED)
    add_text(p, data.get('date',''), size=10.5, color=MUTED)
    add_text(p, ' 申请项目： ', size=10.5, color=MUTED)
    add_text(p, data.get('application_project',''), size=10.5, color=MUTED)
    add_rule(doc)

    bg = data.get('background', {})
    bg['regions_text'] = '、'.join(regions)
    parts = [bg.get('university',''), bg.get('major','')]
    extras = []
    if bg.get('study_length'): extras.append(bg['study_length'])
    if bg.get('graduation_year'): extras.append(f"{bg['graduation_year']}年毕业")
    if extras: parts.append(' · '.join(extras))
    if bg.get('school_note'): parts.append(f"（{bg['school_note']}）")
    bg['school_background'] = ' · '.join([x for x in parts if x])

    add_section_title(doc, '一、学生背景综合评估')
    add_background_table(doc, bg)
    add_rule(doc)

    add_section_title(doc, f"二、{region_text}选校方案")
    include_academic = data.get('include_academic_requirement', True)
    add_programs_table(doc, data.get('programs', []), include_academic)
    add_rule(doc)

    current_year = data.get('date','2026')[:4]
    intake_year = ''.join([c for c in data.get('application_project','') if c.isdigit()])
    intake_year = intake_year[-4:] if len(intake_year) >= 4 else str(int(current_year)+1)

    add_section_title(doc, '四、时间规划建议')
    add_small_heading(doc, f"{current_year}年（专业探究和背景补充）")
    add_timeline_table(doc, data.get('timeline',{}).get('exploration', []))
    add_small_heading(doc, f"{current_year}年（材料准备期）")
    add_timeline_table(doc, data.get('timeline',{}).get('materials', []))
    add_small_heading(doc, f"{intake_year}年（入学前）")
    add_timeline_table(doc, data.get('timeline',{}).get('pre_arrival', []))
    add_rule(doc)
    add_section_title(doc, '五、下一步行动清单')

    # Service page always starts on a new page.
    doc.add_page_break()
    service_img = Path(asset_dir) / 'service-packages.png'
    if service_img.exists():
        pimg = doc.add_paragraph(); pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pimg.add_run().add_picture(str(service_img), width=Cm(17.2))
        format_paragraph(pimg, after=20)
    add_rule(doc)
    p1 = doc.add_paragraph(); format_paragraph(p1, after=8)
    add_text(p1, 'YUSHI留学介绍： ', size=11)
    add_hyperlink(p1, 'https://www.yushiconsulting.cn', 'https://www.yushiconsulting.cn', 11)
    p2 = doc.add_paragraph(); format_paragraph(p2, after=0)
    add_text(p2, '硕士项目介绍： ', size=11)
    add_hyperlink(p2, 'https://masterapply.yushiconsulting.cn', 'https://masterapply.yushiconsulting.cn', 11)

    # Document metadata
    props = doc.core_properties
    props.title = title
    props.subject = 'YUSHI 首咨选校方案'
    props.author = 'YUSHI Consulting'

    doc.save(output_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--input', required=True)
    ap.add_argument('--output', required=True)
    ap.add_argument('--asset-dir', default=str(Path(__file__).resolve().parents[1] / 'assets'))
    args = ap.parse_args()
    with open(args.input, 'r', encoding='utf-8') as f:
        data=json.load(f)
    build(data, args.output, args.asset_dir)

if __name__ == '__main__':
    main()
