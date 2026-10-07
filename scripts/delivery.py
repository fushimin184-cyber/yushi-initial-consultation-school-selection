"""Verified local delivery; does not certify ChatGPT preview or download transport."""
import hashlib
import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from docx import Document


class DeliveryError(RuntimeError):
    pass


def output_path(value):
    path = Path(value).expanduser().absolute()
    # Validate the basename only; parent directories may contain Chinese text.
    if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_-]{0,95}\.docx', path.name):
        raise DeliveryError('Use a short English filename, e.g. Plan_20261008_v2.docx')
    if path.is_symlink():
        raise DeliveryError('Output must not be a symbolic link')
    return path


def file_info(path):
    path = Path(path)
    if not path.is_file() or path.stat().st_size == 0:
        raise DeliveryError(f'Missing or empty delivery file: {path}')
    content = path.read_bytes()  # Also checks actual readability.
    return {'path': str(path.resolve()), 'bytes': len(content),
            'sha256': hashlib.sha256(content).hexdigest()}


def validate_docx(path, expected_xml=None):
    info = file_info(path)
    try:
        with zipfile.ZipFile(path) as archive:
            bad = archive.testzip()
            if bad:
                raise DeliveryError(f'DOCX CRC failure: {bad}')
            required = {'[Content_Types].xml', '_rels/.rels', 'word/document.xml'}
            if not required.issubset(archive.namelist()):
                raise DeliveryError('DOCX is missing required OOXML parts')
            for name in archive.namelist():
                if name.endswith(('.xml', '.rels')):
                    ElementTree.fromstring(archive.read(name))
        reopened = Document(path)
        if expected_xml is not None and reopened._element.xml != expected_xml:
            raise DeliveryError('DOCX content or formatting changed during save')
    except (zipfile.BadZipFile, ElementTree.ParseError, KeyError, ValueError) as exc:
        raise DeliveryError(f'Invalid DOCX: {path}') from exc
    return info


def validate_zip(path, docx_path):
    info = file_info(path)
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() or archive.namelist() != [docx_path.name]:
            raise DeliveryError('Invalid backup ZIP contents')
        content = archive.read(docx_path.name)
        if content != docx_path.read_bytes():
            raise DeliveryError('Backup ZIP does not match the verified DOCX')
        Document(io.BytesIO(content))
    return info


def convert_pdf(docx_path, target):
    executable = shutil.which('soffice') or shutil.which('libreoffice')
    if not executable:
        raise DeliveryError('PDF requested but LibreOffice/soffice is unavailable; DOCX and ZIP remain usable')
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DeliveryError('PDF validation requires pypdf; DOCX and ZIP remain usable') from exc
    with tempfile.TemporaryDirectory(prefix='yushi-pdf-', dir=target.parent) as tmp:
        folder = Path(tmp)
        result = subprocess.run(
            [executable, '-env:UserInstallation=' + (folder / 'profile').as_uri(),
             '--headless', '--convert-to', 'pdf', '--outdir', str(folder), str(docx_path)],
            capture_output=True, text=True, timeout=120, check=False)
        pdf = folder / (docx_path.stem + '.pdf')
        if result.returncode != 0 or not pdf.is_file():
            raise DeliveryError('PDF conversion failed; DOCX and ZIP remain usable')
        file_info(pdf)
        reader = PdfReader(str(pdf), strict=True)
        if reader.is_encrypted or not reader.pages:
            raise DeliveryError('PDF is encrypted or has no pages')
        extracted = ''.join(page.extract_text() or '' for page in reader.pages)
        # Catch missing/substituted Chinese glyphs even when PDF parsing succeeds.
        source = ''.join(node.text or '' for node in
                         Document(docx_path)._element.iter(
                             '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        required_cjk = set(re.findall(r'[\u3400-\u9fff]', source))
        if not required_cjk.issubset(set(extracted)):
            raise DeliveryError('PDF Chinese text verification failed; install the document Chinese fonts and perform visual QA; DOCX and ZIP remain usable')
        pdf.replace(target)
    return file_info(target)


def save_delivery(doc, value, pdf=False):
    target = output_path(value)
    target.parent.mkdir(parents=True, exist_ok=True)
    backup = target.with_suffix('.zip')
    pdf_target = target.with_suffix('.pdf')
    if any(p.is_symlink() for p in (backup, pdf_target)):
        raise DeliveryError('Companion output must not be a symbolic link')
    # Validate staged files before replacing existing deliverables.
    with tempfile.TemporaryDirectory(prefix='yushi-save-', dir=target.parent) as tmp:
        staged = Path(tmp) / target.name
        doc.save(staged)
        validate_docx(staged, doc._element.xml)
        zipped = staged.with_suffix('.zip')
        with zipfile.ZipFile(zipped, 'w', zipfile.ZIP_DEFLATED) as archive:
            archive.write(staged, arcname=target.name)
        validate_zip(zipped, staged)
        # Invalidate old companions before publishing a different DOCX.
        pdf_target.unlink(missing_ok=True)
        staged.replace(target)
        try:
            zipped.replace(backup)
        except OSError:
            backup.unlink(missing_ok=True)
            raise
    manifest = {'docx': validate_docx(target, doc._element.xml),
                'zip': validate_zip(backup, target),
                'validation': 'local integrity verified; visual QA and client download not certified'}
    if pdf:
        try:
            manifest['pdf'] = convert_pdf(target, pdf_target)
        except Exception as exc:
            manifest['pdf_error'] = str(exc)
    return manifest
