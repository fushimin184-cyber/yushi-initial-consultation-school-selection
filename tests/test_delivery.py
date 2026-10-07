import copy
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from build_plan import build
from delivery import DeliveryError, output_path, save_delivery, validate_docx, validate_zip
from docx import Document


class DeliveryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)
        self.data = json.loads((ROOT / 'examples/sample-input.json').read_text())

    def test_build_and_backup_preserve_layout_and_content(self):
        # Baseline OOXML part hashes from commit 8def5d8, fictional sample only.
        expected = json.loads((ROOT / 'tests/baseline-parts.json').read_text())
        new = self.folder / 'nested' / 'Plan_20261008.docx'
        result = build(copy.deepcopy(self.data), new, ROOT / 'assets')
        with zipfile.ZipFile(new) as archive:
            actual = {name: hashlib.sha256(archive.read(name)).hexdigest()
                      for name in archive.namelist()}
        self.assertEqual(expected, actual)
        self.assertEqual(result['docx']['path'], str(new.resolve()))
        self.assertGreater(result['docx']['bytes'], 0)
        self.assertEqual(len(result['docx']['sha256']), 64)
        validate_zip(new.with_suffix('.zip'), new)

    def test_english_basename_unicode_parent(self):
        path = self.folder / '中文目录' / 'Plan_v2.docx'
        save_delivery(Document(), path)
        self.assertTrue(path.exists())
        for name in ('方案.docx', 'Plan space.docx', 'Plan.pdf', '.docx', 'a' * 97 + '.docx'):
            with self.subTest(name=name), self.assertRaises(DeliveryError):
                output_path(self.folder / name)

    def test_missing_empty_corrupt_and_non_docx(self):
        path = self.folder / 'Plan.docx'
        for content in (None, b'', b'not a docx'):
            if content is not None:
                path.write_bytes(content)
            with self.subTest(content=content), self.assertRaises(DeliveryError):
                validate_docx(path)
        with zipfile.ZipFile(path, 'w') as archive:
            archive.writestr('other.txt', 'x')
        with self.assertRaises(DeliveryError):
            validate_docx(path)

    def test_malformed_xml_and_content_mismatch(self):
        path = self.folder / 'Plan.docx'
        Document().save(path)
        with self.assertRaises(DeliveryError):
            validate_docx(path, 'unexpected XML')
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries['word/document.xml'] = b'<broken'
        with zipfile.ZipFile(path, 'w') as archive:
            for name, content in entries.items():
                archive.writestr(name, content)
        with self.assertRaises(DeliveryError):
            validate_docx(path)

    def test_failed_save_leaves_previous_files(self):
        path = self.folder / 'Plan.docx'
        save_delivery(Document(), path)
        before = path.read_bytes(), path.with_suffix('.zip').read_bytes()
        with patch('delivery.validate_docx', side_effect=DeliveryError('forced failure')):
            with self.assertRaises(DeliveryError):
                save_delivery(Document(), path)
        self.assertEqual(before, (path.read_bytes(), path.with_suffix('.zip').read_bytes()))

    def test_crc_failure_is_rejected(self):
        path = self.folder / 'Plan.docx'
        Document().save(path)
        with patch('delivery.zipfile.ZipFile.testzip', return_value='word/document.xml'):
            with self.assertRaises(DeliveryError):
                validate_docx(path)

    def test_pdf_failed_process_timeout_and_missing_output(self):
        import subprocess
        for outcome in (subprocess.CompletedProcess([], 1),
                        subprocess.CompletedProcess([], 0),
                        subprocess.TimeoutExpired('soffice', 120)):
            with self.subTest(outcome=str(outcome)):
                kwargs = {'side_effect': outcome} if isinstance(outcome, Exception) else {'return_value': outcome}
                with patch('delivery.shutil.which', return_value='/fake/soffice'), \
                     patch('delivery.subprocess.run', **kwargs):
                    result = save_delivery(Document(), self.folder / 'Failure.docx', pdf=True)
                self.assertIn('pdf_error', result)
                self.assertNotIn('pdf', result)
                validate_docx(self.folder / 'Failure.docx')
                validate_zip(self.folder / 'Failure.zip', self.folder / 'Failure.docx')

    def test_failed_zip_publish_removes_stale_companions(self):
        path = self.folder / 'Plan.docx'
        save_delivery(Document(), path)
        path.with_suffix('.pdf').write_bytes(b'old pdf')
        original = Path.replace
        def fail_zip(source, target):
            if source.suffix == '.zip':
                raise OSError('forced ZIP publication failure')
            return original(source, target)
        doc = Document()
        doc.add_paragraph('new content')
        with patch.object(Path, 'replace', fail_zip), self.assertRaises(OSError):
            save_delivery(doc, path)
        self.assertEqual(Document(path).paragraphs[0].text, 'new content')
        self.assertFalse(path.with_suffix('.zip').exists())
        self.assertFalse(path.with_suffix('.pdf').exists())

    def test_backup_mismatch(self):
        path = self.folder / 'Plan.docx'
        save_delivery(Document(), path)
        with zipfile.ZipFile(path.with_suffix('.zip'), 'w') as archive:
            archive.writestr(path.name, b'wrong')
        with self.assertRaises(DeliveryError):
            validate_zip(path.with_suffix('.zip'), path)

    def test_pdf_unavailable_reports_failure_keeps_docx_zip(self):
        path = self.folder / 'Plan.docx'
        path.with_suffix('.pdf').write_bytes(b'stale')
        with patch('delivery.shutil.which', return_value=None):
            result = save_delivery(Document(), path, pdf=True)
        self.assertIn('pdf_error', result)
        self.assertNotIn('pdf', result)
        self.assertFalse(path.with_suffix('.pdf').exists())
        validate_docx(path)
        validate_zip(path.with_suffix('.zip'), path)

    def test_default_has_no_pdf_and_symlink_rejected(self):
        path = self.folder / 'Plan.docx'
        result = save_delivery(Document(), path)
        self.assertNotIn('pdf', result)
        self.assertFalse(path.with_suffix('.pdf').exists())
        linked = self.folder / 'Link.docx'
        linked.symlink_to(path)
        with self.assertRaises(DeliveryError):
            save_delivery(Document(), linked)

    @unittest.skipUnless(shutil.which('soffice') or shutil.which('libreoffice'), 'LibreOffice unavailable')
    def test_real_pdf_conversion(self):
        doc = Document()
        doc.add_paragraph('PDF conversion acceptance: fictional sample')
        result = save_delivery(doc, self.folder / 'Plan.docx', pdf=True)
        self.assertNotIn('pdf_error', result, result.get('pdf_error'))
        from pypdf import PdfReader
        self.assertGreater(len(PdfReader(result['pdf']['path']).pages), 0)

    def test_pdf_missing_chinese_text_is_not_delivered(self):
        doc = Document()
        doc.add_paragraph('首咨选校规划')
        from pypdf import PdfWriter
        def fake_conversion(command, **kwargs):
            folder = Path(command[command.index('--outdir') + 1])
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            with (folder / 'Plan.pdf').open('wb') as stream:
                writer.write(stream)
            return subprocess.CompletedProcess(command, 0)
        with patch('delivery.shutil.which', return_value='/fake/soffice'), \
             patch('delivery.subprocess.run', side_effect=fake_conversion):
            result = save_delivery(doc, self.folder / 'Plan.docx', pdf=True)
        self.assertIn('Chinese text verification failed', result['pdf_error'])
        self.assertNotIn('pdf', result)
        self.assertFalse((self.folder / 'Plan.pdf').exists())

    def test_cli_returns_existing_absolute_paths(self):
        result = subprocess.run([sys.executable, str(ROOT / 'scripts/build_plan.py'),
            '--input', str(ROOT / 'examples/sample-input.json'), '--output', str(self.folder / 'CLI.docx')],
            capture_output=True, text=True, check=True)
        manifest = json.loads(result.stdout)
        for kind in ('docx', 'zip'):
            self.assertTrue(Path(manifest[kind]['path']).is_file())
            self.assertTrue(Path(manifest[kind]['path']).is_absolute())


if __name__ == '__main__':
    unittest.main()
